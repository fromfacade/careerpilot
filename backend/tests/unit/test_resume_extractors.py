"""Unit tests for app.resume_parsing.extractors.

No binary fixture files are checked into the repo: the DOCX fixture is
built in-memory with `python-docx` (the exact library the extractor
itself uses to read it), and the PDF fixture is a hand-assembled, byte-
exact-offset minimal PDF -- a standard technique for generating a real,
valid, parseable PDF without a drawing/layout library. Both avoid adding
a test-only dependency just to produce fixtures.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.core.exceptions import ResumeExtractionError
from app.models.enums import ResumeFileFormat
from app.resume_parsing.extractors import extract_text


def _build_minimal_pdf(lines: list[str]) -> bytes:
    """Hand-assemble a minimal, single-page, valid PDF whose content
    stream shows each of `lines` as real extractable text, with exact
    byte offsets in the xref table (no reliance on pypdf's malformed-PDF
    recovery path)."""
    content_ops = ["BT", "/F1 12 Tf", "72 712 Td"]
    for i, line in enumerate(lines):
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        if i > 0:
            content_ops.append("0 -14 Td")
        content_ops.append(f"({escaped}) Tj")
    content_ops.append("ET")
    stream_body = "\n".join(content_ops).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream_body)).encode("ascii") + b" >>\nstream\n" + stream_body + b"\nendstream",
    ]

    buffer = io.BytesIO()
    buffer.write(b"%PDF-1.4\n")
    offsets = [0]  # object 0 is the free-list head, offset unused
    for i, body in enumerate(objects, start=1):
        offsets.append(buffer.tell())
        buffer.write(f"{i} 0 obj\n".encode("ascii"))
        buffer.write(body)
        buffer.write(b"\nendobj\n")

    xref_offset = buffer.tell()
    n = len(objects) + 1
    buffer.write(f"xref\n0 {n}\n".encode("ascii"))
    buffer.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        buffer.write(f"{off:010d} 00000 n \n".encode("ascii"))
    buffer.write(
        f"trailer\n<< /Size {n} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii")
    )
    return buffer.getvalue()


def test_extract_pdf_text_returns_real_lines():
    pdf_bytes = _build_minimal_pdf(["Jane Doe", "Software Engineer"])
    text = extract_text(pdf_bytes, ResumeFileFormat.PDF)
    assert "Jane Doe" in text
    assert "Software Engineer" in text


def test_extract_pdf_text_rejects_corrupt_bytes():
    with pytest.raises(ResumeExtractionError):
        extract_text(b"not a real pdf at all, just some bytes", ResumeFileFormat.PDF)


def _build_docx_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_docx_text_returns_paragraphs():
    docx_bytes = _build_docx_bytes(["Jane Doe", "EDUCATION", "UC Santa Cruz - B.S. Computer Science"])
    text = extract_text(docx_bytes, ResumeFileFormat.DOCX)
    assert "Jane Doe" in text
    assert "EDUCATION" in text
    assert "UC Santa Cruz - B.S. Computer Science" in text


def test_extract_docx_text_includes_table_cells():
    document = Document()
    document.add_paragraph("Skills")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "TypeScript"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text(buffer.getvalue(), ResumeFileFormat.DOCX)
    assert "Python" in text
    assert "TypeScript" in text


def test_extract_docx_text_rejects_corrupt_bytes():
    with pytest.raises(ResumeExtractionError):
        extract_text(b"definitely not a zip/docx file", ResumeFileFormat.DOCX)


def test_extract_pdf_text_rejects_empty_text_pdf():
    # A structurally valid PDF with an empty content stream has nothing to
    # extract -- this must be a clear ResumeExtractionError, not an empty
    # string silently treated as success.
    pdf_bytes = _build_minimal_pdf([])
    with pytest.raises(ResumeExtractionError):
        extract_text(pdf_bytes, ResumeFileFormat.PDF)
