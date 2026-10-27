from pathlib import Path

from app.services.resume_export import export_docx, export_pdf
from app.services.resume_tailor import flatten_tailored_lines, tailor_resume, visible_rows
from docx import Document
from pypdf import PdfReader

from tests.resume_export_fixtures import (
    AI_ANALYSIS,
    BACKEND_ANALYSIS,
    FULLSTACK_ANALYSIS,
    Candidate,
    realistic_facts,
)


def _render_pages(pdf_path: Path, dest: Path) -> list[Path]:
    import pypdfium2 as pdfium

    dest.mkdir(parents=True, exist_ok=True)
    document = pdfium.PdfDocument(str(pdf_path))
    paths = []
    for index, page in enumerate(document, start=1):
        image = page.render(scale=2).to_pil()
        out = dest / f"page-{index}.png"
        image.save(out)
        paths.append(out)
    return paths


def _ink_ratio(image_path: Path) -> float:
    from PIL import Image

    image = Image.open(image_path).convert("L")
    pixels = list(image.getdata())
    ink = sum(1 for pixel in pixels if pixel < 245)
    return ink / max(len(pixels), 1)


def test_backend_and_fullstack_select_different_projects():
    facts = realistic_facts()
    original = [dict(fact.fact_value) for fact in facts]
    backend = tailor_resume(Candidate(), facts, BACKEND_ANALYSIS)
    fullstack = tailor_resume(Candidate(), facts, FULLSTACK_ANALYSIS)
    ai = tailor_resume(Candidate(), facts, AI_ANALYSIS)
    backend_names = {row["title"] for row in visible_rows(backend, "projects")}
    fullstack_names = {row["title"] for row in visible_rows(fullstack, "projects")}
    ai_names = {row["title"] for row in visible_rows(ai, "projects")}
    assert {"TaskFlow", "GlucoRelay"} <= backend_names
    assert {"PC Build Visual Builder", "Spritz"} <= fullstack_names
    assert "GlucoRelay" in ai_names and "Spritz" in ai_names
    assert 2 <= len(backend_names) <= 3
    assert all(row.get("selection_reason") for row in visible_rows(backend, "projects"))
    assert all(len((row["detail"].get("included_bullet_indexes") or [])) <= 3 for row in visible_rows(backend, "projects"))
    assert [fact.fact_value for fact in facts] == original


def test_skills_stay_grouped_and_job_relevant_first():
    tailored = tailor_resume(Candidate(), realistic_facts(), BACKEND_ANALYSIS)
    skills = visible_rows(tailored, "skills")
    titles = [row["title"] for row in skills]
    assert titles[:5] == [
        "Programming Languages",
        "Backend / APIs",
        "Frontend",
        "Databases / Infrastructure",
        "Tools / Platforms",
    ]
    flatten = "\n".join(flatten_tailored_lines(tailored))
    assert "Programming Languages: Python" in flatten
    assert "Backend / APIs:" in flatten
    assert "Languages: English, Spanish" in flatten
    backend_items = next(row["detail"]["items"] for row in skills if row["title"] == "Backend / APIs")
    assert backend_items[0] in {"FastAPI", "REST APIs"}


def test_manual_include_override_and_export_consistency(tmp_path: Path):
    tailored = tailor_resume(Candidate(), realistic_facts(), BACKEND_ANALYSIS)
    hidden = next(row for row in tailored["projects"] if row.get("included") is False)
    assert hidden["included"] is False
    hidden["included"] = True
    hidden["selection_reason"] = "Manually included."
    tailored["presentation"]["page_budget"] = "manual"
    pdf_path = export_pdf(tailored, tmp_path / "manual.pdf")
    docx_path = export_docx(tailored, tmp_path / "manual.docx")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages)
    docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    for text in (pdf_text, docx_text):
        assert "Spritz" in text
        assert "TaskFlow" in text
        assert "SHPE" in text
        assert text.count("Participated in professional development and peer mentorship initiatives supporting students in STEM.") == 1


def test_realistic_backend_export_is_one_page_and_renders(tmp_path: Path):
    tailored = tailor_resume(Candidate(), realistic_facts(), BACKEND_ANALYSIS)
    pdf_path = export_pdf(tailored, tmp_path / "backend.pdf")
    pages = [page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages]
    assert len(pages) == 1
    text = pages[0]
    assert "Alex Rivera" in text
    assert "github.com/alexrivera" in text
    assert "alexrivera.dev" in text
    assert "TaskFlow" in text
    assert "GlucoRelay" in text
    assert "Programming Languages:" in text
    assert "SHPE" in text
    assert text.count("SHPE") == 1
    assert "Dates:" not in text
    assert "? Built" not in text
    images = _render_pages(pdf_path, tmp_path / "backend-pages")
    assert len(images) == 1
    assert _ink_ratio(images[0]) > 0.04
    assert tailored["quality"]["ok"] is True
    assert tailored["quality"]["page_count"] == 1


def test_fullstack_export_is_one_page(tmp_path: Path):
    tailored = tailor_resume(Candidate(), realistic_facts(), FULLSTACK_ANALYSIS)
    pdf_path = export_pdf(tailored, tmp_path / "fullstack.pdf")
    pages = [page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages]
    assert len(pages) == 1
    text = pages[0]
    assert "PC Build Visual Builder" in text
    assert "Spritz" in text
    images = _render_pages(pdf_path, tmp_path / "fullstack-pages")
    assert len(images) == 1
    assert _ink_ratio(images[0]) > 0.04


def test_three_projects_do_not_push_skills_off_page_one(tmp_path: Path):
    """Reproduce tailored-resume-4.pdf: three projects left Skills and SHPE on a sparse page 2."""
    facts = realistic_facts()
    original = [dict(fact.fact_value) for fact in facts]
    tailored = tailor_resume(Candidate(), facts, BACKEND_ANALYSIS)
    extra = " Coordinated deployment notes, structured logging, and production error handling with the rest of the team."
    for row in tailored["projects"]:
        if row["title"] in {"TaskFlow", "GlucoRelay", "Spritz"}:
            row["included"] = True
            detail = row["detail"] if isinstance(row["detail"], dict) else {}
            bullets = [f"{item}{extra}" for item in (detail.get("bullets") or [])]
            detail["bullets"] = bullets
            detail["included_bullet_indexes"] = list(range(min(3, len(bullets))))
            row["detail"] = detail
    for row in tailored["achievements"]:
        row["included"] = True
    overloaded = {**tailored, "presentation": {**tailored.get("presentation", {}), "page_budget": "manual"}}
    from app.services.resume_export import _pdf_page_stats, _render_pdf_document

    overflow_pages, overflow_texts = _pdf_page_stats(_render_pdf_document(overloaded))
    assert overflow_pages == 2
    assert "Programming Languages:" in overflow_texts[1] or "SKILLS" in overflow_texts[1]
    tailored["presentation"]["page_budget"] = "auto"
    pdf_path = export_pdf(tailored, tmp_path / "overflow.pdf")
    pages = [page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages]
    assert len(pages) == 1
    assert "SKILLS" in pages[0]
    assert "Programming Languages:" in pages[0]
    assert "TaskFlow" in pages[0]
    included_projects = [row["title"] for row in visible_rows(tailored, "projects")]
    assert 2 <= len(included_projects) <= 3
    if len(included_projects) < 3:
        assert tailored["presentation"].get("tradeoff")
    images = _render_pages(pdf_path, tmp_path / "overflow-pages")
    assert len(images) == 1
    assert _ink_ratio(images[0]) > 0.04
    assert [fact.fact_value for fact in facts] == original


def test_specific_mechanism_outranks_generic_architecture_bullet():
    from app.services.resume_tailor import strongest_bullet_indexes
    from app.services.resume_tailor import _tokens

    taskflow = next(fact for fact in realistic_facts() if fact.fact_key == "TaskFlow")
    bullets = taskflow.fact_value["bullets"]
    target = _tokens("FastAPI PostgreSQL Docker REST APIs SQL background job processing")
    indexes = strongest_bullet_indexes(
        bullets,
        target,
        2,
        technologies=taskflow.fact_value["technologies"],
    )
    chosen = " ".join(bullets[index] for index in indexes)
    assert "concurrency-safe" in chosen
    assert "clear split" not in chosen


def test_contact_line_uses_verified_links_only():
    present = tailor_resume(Candidate(), realistic_facts(), BACKEND_ANALYSIS)
    contact = " ".join(flatten_tailored_lines(present)[:2])
    assert "github.com/alexrivera" in contact
    assert "alexrivera.dev" in contact
    assert "alex.rivera@example.com" in contact
    missing = Candidate()
    missing.github_url = None
    missing.portfolio_url = None
    hidden = tailor_resume(missing, realistic_facts(), BACKEND_ANALYSIS)
    contact = " ".join(flatten_tailored_lines(hidden)[:2])
    assert "github.com" not in contact
    assert "alexrivera.dev" not in contact
    assert "http" not in contact


def test_internship_resume_keeps_three_projects_with_two_bullets(tmp_path: Path):
    facts = realistic_facts()
    original = [dict(fact.fact_value) for fact in facts]
    tailored = tailor_resume(Candidate(), facts, BACKEND_ANALYSIS)
    pdf_path = export_pdf(tailored, tmp_path / "three.pdf")
    pages = [page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages]
    included = visible_rows(tailored, "projects")
    assert len(pages) == 1
    assert len(included) == 3
    assert "SKILLS" in pages[0]
    assert "Programming Languages:" in pages[0]
    assert tailored["presentation"]["body_pt"] >= 9
    for row in included:
        indexes = row["detail"].get("included_bullet_indexes") or []
        assert 1 <= len(indexes) <= 2
        assert row["title"] in pages[0]
    for row in visible_rows(tailored, "work_experience"):
        indexes = row["detail"].get("included_bullet_indexes") or []
        assert 1 <= len(indexes) <= 2
    images = _render_pages(pdf_path, tmp_path / "three-pages")
    assert len(images) == 1
    assert _ink_ratio(images[0]) > 0.04
    assert [fact.fact_value for fact in facts] == original
