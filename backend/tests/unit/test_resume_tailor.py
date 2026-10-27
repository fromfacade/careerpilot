from pathlib import Path

from app.models.enums import FactType
from app.services.resume_export import export_docx, export_pdf
from app.services.resume_tailor import (
    assess_resume_quality,
    flatten_tailored_lines,
    structured_entry_lines,
    tailor_resume,
    visible_rows,
)
from docx import Document
from pypdf import PdfReader


class _Fact:
    def __init__(self, id, fact_type, fact_key, fact_value, identity_key=None):
        self.id = id
        self.fact_type = fact_type
        self.fact_key = fact_key
        self.fact_value = fact_value
        self.identity_key = identity_key or f"{fact_type.value}|{fact_key.casefold()}"


class _Candidate:
    full_name = "Jane Doe"
    email = "jane@example.com"
    phone = "(555) 010-0100"
    current_location = "California"
    school = "UCSC"
    linkedin_url = None
    github_url = "github.com/janedoe"
    portfolio_url = None


def _intern_facts():
    return [
        _Fact(
            1,
            FactType.EDUCATION,
            "UCSC",
            {
                "institution": "UC Santa Cruz",
                "degree": "B.S. Computer Science",
                "dates_text": "Expected Graduation: December 2027",
                "gpa": "3.50/4.00",
                "coursework": "Data Structures, Operating Systems, Databases",
            },
            identity_key="education|ucsc",
        ),
        _Fact(
            2,
            FactType.EDUCATION,
            "University of California, Santa Cruz",
            {
                "institution": "University of California, Santa Cruz",
                "degree": "B.S. Computer Science",
                "location": "Santa Cruz, CA",
            },
            identity_key="education|university of california, santa cruz",
        ),
        _Fact(
            3,
            FactType.WORK_EXPERIENCE,
            "Northwind",
            {
                "employer": "Northwind Tutoring",
                "position": "Office Assistant",
                "location": "Springfield, ST",
                "dates_text": "September 2024 - Present",
                "bullets": [
                    "Built and automated a utilization-tracking system used by staff.",
                    "Created reusable spreadsheet logic and workflow documentation.",
                ],
            },
            identity_key="work_experience|northwind",
        ),
        _Fact(
            4,
            FactType.WORK_EXPERIENCE,
            "Northwind Tutoring Services",
            {
                "employer": "Northwind Tutoring Services",
                "position": "Student Office Assistant",
                "dates_text": "September 2024 - Present",
                "bullets": ["Built and automated a utilization-tracking system used by staff."],
            },
            identity_key="work_experience|northwind tutoring services",
        ),
        _Fact(
            5,
            FactType.WORK_EXPERIENCE,
            "Contoso",
            {
                "employer": "Contoso Student Life",
                "position": "Tech Crew Member",
                "location": "Springfield, ST",
                "dates_text": "September 2024 - June 2025",
                "bullets": ["Configured and operated A/V systems for events with 150+ attendees."],
            },
        ),
        _Fact(
            6,
            FactType.PROJECT,
            "Campus Site",
            {"name": "Campus Site", "technologies": ["HTML"], "bullets": ["Published club pages."]},
        ),
        _Fact(
            7,
            FactType.PROJECT,
            "Alpha Queue",
            {
                "name": "Alpha Queue",
                "technologies": ["FastAPI", "Redis", "PostgreSQL"],
                "bullets": ["Queued jobs.", "Implemented JWT access."],
            },
        ),
        _Fact(
            8,
            FactType.PROJECT,
            "Alpha Queue Platform",
            {"name": "Alpha Queue — Backend Job Processing", "technologies": ["FastAPI"]},
        ),
        _Fact(
            9,
            FactType.PROJECT,
            "Beta Relay",
            {"name": "Beta Relay", "technologies": ["FastAPI", "Gemini API"], "bullets": ["Built multilingual check-ins."]},
        ),
        _Fact(
            10,
            FactType.PROJECT,
            "Gamma Scent",
            {"name": "Gamma Scent", "technologies": ["Node.js", "Express"], "bullets": ["Shipped a hackathon MVP."]},
        ),
        _Fact(
            11,
            FactType.PROJECT,
            "Delta Builder",
            {"name": "Delta Builder", "technologies": ["React", "Next.js"], "bullets": ["Filtered compatible PC parts."]},
        ),
        _Fact(12, FactType.SKILL, "Languages", {"items": ["Python", "TypeScript", "SQL"]}),
        _Fact(13, FactType.SKILL, "Backend", {"items": ["FastAPI", "Python", "PostgreSQL"]}),
        _Fact(
            14,
            FactType.ACHIEVEMENT,
            "Tracker",
            {"text": "Built and automated a utilization-tracking system used by staff."},
        ),
        _Fact(
            15,
            FactType.ACHIEVEMENT,
            "SEE",
            {"text": "Participated in professional development and peer mentorship.", "organization": "SEE"},
        ),
    ]


def test_tailor_reorders_relevant_projects_and_does_not_invent():
    facts = [
        _Fact(1, FactType.PROJECT, "Campus Site", {"name": "Campus Site", "technologies": ["HTML"]}),
        _Fact(2, FactType.PROJECT, "Alpha Queue", {"name": "Alpha Queue", "technologies": ["FastAPI", "Redis"]}),
        _Fact(3, FactType.SKILL, "Languages", {"items": ["Python"]}),
    ]
    analysis = {
        "summary": "Backend engineer using FastAPI",
        "technologies": [{"text": "fastapi", "evidence": "stated"}],
        "required": [{"text": "Python", "evidence": "stated"}],
    }
    tailored = tailor_resume(_Candidate(), facts, analysis)
    assert [row["title"] for row in tailored["projects"] if row["included"]][0] == "Alpha Queue"
    blob = str(tailored)
    assert "NASA" not in blob
    assert "Alpha Queue" in blob
    assert tailored["source_fact_ids"] == [1, 2, 3]


def test_tailor_dedupes_semantic_duplicates_without_dropping_distinct_roles():
    facts = _intern_facts()
    original_values = [dict(fact.fact_value) for fact in facts]
    tailored = tailor_resume(
        _Candidate(),
        facts,
        {"summary": "Software engineering intern using FastAPI and Python", "technologies": [{"text": "FastAPI"}]},
    )
    education = visible_rows(tailored, "education")
    work = visible_rows(tailored, "work_experience")
    projects = visible_rows(tailored, "projects")
    assert len(education) == 1
    assert "Santa Cruz" in str(education[0]["detail"])
    assert {row["detail"]["employer"] for row in work} == {"Northwind Tutoring", "Contoso Student Life"}
    assert [row["title"] for row in projects][0] == "Alpha Queue"
    assert {row["title"] for row in projects} == {"Alpha Queue", "Beta Relay", "Gamma Scent", "Delta Builder"}
    assert all(row["included"] is False for row in tailored["projects"] if row["title"] == "Campus Site")
    skills = visible_rows(tailored, "skills")
    assert len(skills) == 1
    assert skills[0]["detail"]["items"].count("Python") == 1
    assert any(row["included"] is False for row in tailored["achievements"] if row["title"] == "Tracker")
    assert any(row["included"] is True for row in tailored["achievements"] if row["title"] == "SEE")
    assert [fact.fact_value for fact in facts] == original_values
    assert set(tailored["source_fact_ids"]) == {fact.id for fact in facts}


def test_export_preserves_text_order(tmp_path: Path):
    document = {
        "header": {"full_name": "Jane Doe", "email": "jane@example.com"},
        "education": [{"title": "UCSC", "detail": {"school": "UCSC"}, "included": True}],
        "projects": [{"title": "Alpha Queue", "detail": {"name": "Alpha Queue", "bullets": ["Queued jobs."]}, "included": True}],
        "work_experience": [],
        "skills": [],
        "achievements": [],
    }
    docx_path = export_docx(document, tmp_path / "t.docx")
    pdf_path = export_pdf(document, tmp_path / "t.pdf")
    docx_text = "\n".join(p.text for p in Document(docx_path).paragraphs)
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages)
    for text in (docx_text, pdf_text):
        assert text.casefold().index("jane doe") < text.casefold().index("education") < text.casefold().index("ucsc")
        assert text.casefold().index("projects") < text.casefold().index("alpha queue")
    assert flatten_tailored_lines(document)[0] == "Jane Doe"
    for text in (docx_text, pdf_text, "\n".join(flatten_tailored_lines(document))):
        assert "{" not in text
        assert "Queued jobs." in text
        assert "'name':" not in text
        assert "Dates:" not in text
        assert "Subtitle:" not in text


def test_structured_entry_renders_work_and_project_fields():
    work = structured_entry_lines(
        {
            "title": "Northwind",
            "detail": {
                "employer": "Northwind Tutoring",
                "position": "Office Assistant",
                "dates_text": "2024 – Present",
                "bullets": ["Built a tracker.", "Documented workflows."],
                "review_reasons": ["ignore this"],
            },
        },
        "work_experience",
    )
    assert work[0] == "Office Assistant - Northwind Tutoring"
    assert "2024 - Present" in work
    assert "Dates:" not in "\n".join(work)
    assert "- Built a tracker." in work
    assert "review_reasons" not in " ".join(work)
    assert "{" not in "".join(work)

    project = structured_entry_lines(
        {
            "title": "Alpha Queue",
            "detail": {"name": "Alpha Queue", "technologies": ["FastAPI", "Redis"], "bullets": ["Queued jobs."]},
        },
        "projects",
    )
    assert project[0] == "Alpha Queue"
    assert "FastAPI, Redis" in project
    assert "Technologies:" not in "\n".join(project)
    assert "- Queued jobs." in project


def test_hidden_rows_do_not_render():
    hidden = structured_entry_lines(
        {"title": "Campus Site", "detail": {"name": "Campus Site"}, "included": False},
        "projects",
    )
    assert hidden == []


def test_pdf_export_keeps_extractable_text_on_long_resume(tmp_path: Path):
    bullets = [f"Delivered feature number {index} for the campus platform." for index in range(1, 40)]
    document = {
        "header": {"full_name": "Alex Rivera", "email": "alex.rivera@example.com", "location": "Santa Cruz, CA"},
        "education": [{"title": "UCSC", "detail": {"institution": "UC Santa Cruz", "degree": "B.S. Computer Science"}}],
        "work_experience": [
            {
                "title": "Intern",
                "detail": {"employer": "Northwind Labs", "position": "Software Intern", "bullets": bullets[:12]},
            }
        ],
        "projects": [
            {
                "title": "TaskFlow",
                "detail": {"name": "TaskFlow", "technologies": ["FastAPI", "PostgreSQL"], "bullets": bullets[12:24]},
            }
        ],
        "skills": [{"title": "Languages", "detail": {"items": ["Python", "TypeScript"]}}],
        "achievements": [],
    }
    pdf_path = export_pdf(document, tmp_path / "long.pdf")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages)
    assert "Alex Rivera" in pdf_text
    assert "EDUCATION" in pdf_text
    assert "TaskFlow" in pdf_text
    assert "Delivered feature number 1" in pdf_text
    assert "Delivered feature number 24" in pdf_text
    assert "{" not in pdf_text
    assert "? Delivered" not in pdf_text
    pages = PdfReader(str(pdf_path)).pages
    assert all((page.extract_text() or "").strip() for page in pages)
    assert len(pages) <= 2


def test_intern_fixture_export_is_one_page_and_ascii_safe(tmp_path: Path):
    tailored = tailor_resume(
        _Candidate(),
        _intern_facts(),
        {"summary": "Software engineering intern", "technologies": [{"text": "FastAPI"}, {"text": "Python"}]},
    )
    pdf_path = export_pdf(tailored, tmp_path / "intern.pdf")
    docx_path = export_docx(tailored, tmp_path / "intern.docx")
    pages = [page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages]
    pdf_text = "\n".join(pages)
    docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    assert len(pages) == 1
    assert pages[0].strip()
    for text in (pdf_text, docx_text, "\n".join(flatten_tailored_lines(tailored))):
        assert text.count("Jane Doe") == 1
        assert text.count("UC Santa Cruz") == 1
        assert text.count("Office Assistant") == 1
        assert text.count("Alpha Queue") == 1
        assert "Dates:" not in text
        assert "Subtitle:" not in text
        assert "Skills:" not in text
        assert "Technologies:" not in text
        assert "? Built" not in text
        assert "? Queued" not in text
        assert "review_reasons" not in text
    flatten = "\n".join(flatten_tailored_lines(tailored))
    assert "- Queued jobs." in flatten
    assert flatten.count("Participated in professional development and peer mentorship.") == 1
    assert tailored["quality"]["ok"] is True
    assert tailored["quality"]["page_count"] == 1


def test_export_fails_when_required_content_is_missing(tmp_path: Path):
    from app.core.exceptions import ValidationError
    import pytest

    document = {"header": {"full_name": "", "email": ""}, "education": [], "work_experience": [], "projects": [], "skills": [], "achievements": []}
    with pytest.raises(ValidationError, match="no content"):
        export_pdf(document, tmp_path / "empty.pdf")


def test_quality_report_rejects_labels_and_duplicates():
    document = {
        "header": {"full_name": "Jane Doe", "email": "jane@example.com"},
        "education": [
            {"title": "UCSC", "detail": {"institution": "UCSC"}, "included": True},
            {"title": "UCSC", "detail": {"institution": "UCSC"}, "included": True},
        ],
        "work_experience": [],
        "projects": [],
        "skills": [{"title": "Skills", "detail": {"items": ["Python"]}}],
        "achievements": [],
    }
    labeled = flatten_tailored_lines(document)
    assert labeled.count("UCSC") == 2
    quality = assess_resume_quality(document, extracted="Jane Doe\nDates: 2024\n? Built a tracker.", page_count=4, empty_pages=1)
    assert quality["ok"] is False
    assert any("Duplicate" in issue for issue in quality["issues"])
    assert any("blank" in issue for issue in quality["issues"])
    assert any("replacement" in issue for issue in quality["issues"])
    assert any("pages" in issue for issue in quality["issues"])
