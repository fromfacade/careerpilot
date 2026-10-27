"""Shared resume fixtures for parser, extractor, and API tests.

Nothing binary is checked in. Real PDF and DOCX files are generated in
memory:

* PDF: a hand-assembled, valid, multi-page PDF (WinAnsi-encoded Helvetica so
  the bullet, en dash, and em dash glyphs extract exactly like a real
  export) -- no layout library needed.
* DOCX: built with `python-docx`, the same library the extractor reads with.
  Bullets use Word's real "List Bullet" style, so there is *no* bullet glyph
  in the text -- exactly the situation real DOCX resumes present.

The people, employers, and projects below are fictional. The layout mirrors
a common single-column resume: `Employer — Position` then
`Location | Dates`, project title then a `Technologies:` line, then bullets.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from docx import Document

# --------------------------------------------------------------------------
# Content model
# --------------------------------------------------------------------------
# A section body is a list of (kind, text):
#   "text"   - an ordinary line (a header line, metadata line, ...)
#   "bullet" - a bullet; PDF gets a glyph, DOCX gets the List Bullet style
#   "wrap"   - the continuation of the previous bullet. PDF: its own physical
#              line (as a real wrapped line is); DOCX: merged into the bullet
#              paragraph (a Word paragraph has no hard wraps).
Item = tuple[str, str]


def t(text: str) -> Item:
    return ("text", text)


def b(text: str) -> Item:
    return ("bullet", text)


def w(text: str) -> Item:
    return ("wrap", text)


@dataclass
class Section:
    heading: str
    items: list[Item] = field(default_factory=list)


HEADER_LINES = ["Jane Doe", "jane.doe@example.com | (555) 010-0100 | github.com/janedoe"]

EDUCATION = Section(
    "EDUCATION",
    [
        t("Example State University — Springfield, ST"),
        t("B.S. Computer Science | Expected Graduation: December 2027 | GPA: 3.50/4.00"),
        t("Relevant Coursework: Data Structures & Algorithms, Operating Systems, Databases"),
    ],
)

SKILLS = Section(
    "TECHNICAL SKILLS :",
    [
        t("Programming Languages: Python, JavaScript, TypeScript, SQL, C++"),
        t("Backend / APIs: FastAPI, Node.js, Express, REST APIs, SQLAlchemy, Pydantic, JWT"),
        t("Frontend: React, Next.js"),
        t("Databases / Infrastructure: PostgreSQL, Docker, Docker Compose, Linux/Unix, Uvicorn"),
        t("Tools / Platforms: Git/GitHub, Vercel, Google Apps Script, Google Gemini API, Gemma"),
        t("Languages: English, Spanish"),
    ],
)

PROJECTS = Section(
    "PROJECTS :",
    [
        t("Alpha Queue — Backend Job Processing & Authentication Platform"),
        t("Technologies: Python, FastAPI, PostgreSQL, SQLAlchemy, Docker, JWT, REST APIs"),
        b("Built a Python/FastAPI backend service with PostgreSQL persistence for asynchronous job creation, status tracking,"),
        w("and background processing."),
        b("Implemented concurrency-safe job claiming so multiple workers can process queued tasks without duplicate execution."),
        b("Designed REST APIs and SQLAlchemy data models with clear separation between database, service, and API layers."),
        b("Developed secure password hashing, JWT-based access, and email verification workflows."),
        t("Beta Relay — AI-Powered Safety Assistant"),
        t("Technologies: Python, FastAPI, Pydantic, Gemini API, Gemma, REST APIs"),
        b("Built a FastAPI backend for multilingual voice check-ins, converting natural-language responses into structured states."),
        b("Integrated Google Gemma through the Gemini API to extract structured information while keeping safety-critical"),
        w("transitions controlled by deterministic backend logic."),
        b("Designed an idempotent state machine with timed rechecks and caregiver handoff endpoints."),
        t("Gamma Scent — Affordable Alternatives to Luxury Fragrances"),
        t("Technologies: Node.js, Express, Puppeteer, Gemini, PostgreSQL, Docker"),
        b("Built backend workflows in Node.js/Express using Puppeteer and Gemini to analyze fragrance profiles."),
        b("Designed and queried a Dockerized PostgreSQL schema to store fragrance notes and return matches."),
        b("Shipped a working MVP in 36 hours with a 4-person team during a hackathon."),
        t("Delta Builder — Full-Stack PC Compatibility Tool"),
        t("Technologies: React, Next.js, TypeScript, Node.js, Express, PostgreSQL, Docker"),
        b("Developed a full-stack web application with a React/Next.js frontend and Node.js/Express backend."),
        b("Implemented constraint-driven component selection that dynamically filters CPUs, GPUs, and memory."),
        b("Designed a PostgreSQL schema for component specifications and containerized backend services."),
    ],
)

EXPERIENCE = Section(
    "EXPERIENCE :",
    [
        t("Northwind Tutoring Services — Student Office Assistant"),
        t("Springfield, ST | September 2024 – Present"),
        b("Built and automated a utilization-tracking system used by staff to identify underused tutoring services."),
        b("Analyzed attendance, cancellation, and scheduling trends to surface underutilized sessions and"),
        w("support data-informed changes to tutoring coverage."),
        b("Created reusable spreadsheet logic and workflow documentation to reduce manual tracking."),
        t("Contoso Student Life — Tech Crew Member"),
        t("Springfield, ST | September 2024 – June 2025"),
        b("Configured and operated A/V systems for events with 150+ attendees, troubleshooting failures under time constraints."),
        b("Documented technical best practices and contributed process improvements during weekly team syncs."),
    ],
)

ACTIVITIES = Section(
    "ACTIVITIES :",
    [
        t("Society of Example Engineers (SEE), State University Chapter"),
        b("Participated in professional development and peer mentorship initiatives supporting students in STEM."),
    ],
)

STANDARD_SECTIONS = [EDUCATION, SKILLS, PROJECTS, EXPERIENCE, ACTIVITIES]

# The exact record counts every layout variant of the standard resume must yield.
EXPECTED_COUNTS = {
    "education": 1,
    "work_experience": 2,
    "project": 4,
    "skill": 6,
    "achievement": 1,
}
EXPECTED_PROJECT_NAMES = ["Alpha Queue", "Beta Relay", "Gamma Scent", "Delta Builder"]
EXPECTED_SKILL_CATEGORIES = [
    "Programming Languages",
    "Backend / APIs",
    "Frontend",
    "Databases / Infrastructure",
    "Tools / Platforms",
    "Languages",
]


def reordered_sections() -> list[Section]:
    """Same content, different order and synonym headings."""
    return [
        Section("WORK EXPERIENCE", EXPERIENCE.items),
        Section("SELECTED PROJECTS", PROJECTS.items),
        Section("SKILLS", SKILLS.items),
        Section("EDUCATION", EDUCATION.items),
        Section("LEADERSHIP", ACTIVITIES.items),
    ]


# --------------------------------------------------------------------------
# Rendering
# --------------------------------------------------------------------------
def pdf_lines(sections: list[Section], header: list[str] | None = None, blank_between: bool = True) -> list[str]:
    lines = list(HEADER_LINES if header is None else header)
    for section in sections:
        if blank_between:
            lines.append("")
        lines.append(section.heading)
        for kind, text in section.items:
            lines.append(f"• {text}" if kind == "bullet" else text)
    return lines


def docx_paragraphs(sections: list[Section], header: list[str] | None = None) -> list[tuple[bool, str]]:
    """`(is_bullet, text)` pairs; wrapped continuations merge into the bullet."""
    paragraphs: list[tuple[bool, str]] = [(False, line) for line in (HEADER_LINES if header is None else header)]
    for section in sections:
        paragraphs.append((False, section.heading))
        for kind, text in section.items:
            if kind == "wrap" and paragraphs and paragraphs[-1][0]:
                paragraphs[-1] = (True, f"{paragraphs[-1][1]} {text}")
            else:
                paragraphs.append((kind == "bullet", text))
    return paragraphs


def build_pdf(lines: list[str], lines_per_page: int = 52) -> bytes:
    """A valid multi-page PDF with each of `lines` as real extractable text."""

    def encode(text: str) -> bytes:
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        return escaped.encode("cp1252")

    pages = [lines[i : i + lines_per_page] for i in range(0, max(len(lines), 1), lines_per_page)] or [[]]
    objects: list[bytes] = []  # filled so object N is objects[N-1]
    page_count = len(pages)
    font_obj = 3 + 2 * page_count  # after catalog, pages, and (page, content) pairs
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(page_count))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode("ascii"))
    for index, page_lines in enumerate(pages):
        content_obj = 4 + 2 * index
        ops = [b"BT", b"/F1 10 Tf", b"54 750 Td"]
        for i, line in enumerate(page_lines):
            if i > 0:
                ops.append(b"0 -13.5 Td")
            if line:
                ops.append(b"(" + encode(line) + b") Tj")
        ops.append(b"ET")
        stream = b"\n".join(ops)
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 {font_obj} 0 R >> >> "
                f"/MediaBox [0 0 612 792] /Contents {content_obj} 0 R >>"
            ).encode("ascii")
        )
        objects.append(b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")

    buffer = io.BytesIO()
    buffer.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(buffer.tell())
        buffer.write(f"{number} 0 obj\n".encode("ascii"))
        buffer.write(body)
        buffer.write(b"\nendobj\n")
    xref_offset = buffer.tell()
    size = len(objects) + 1
    buffer.write(f"xref\n0 {size}\n".encode("ascii"))
    buffer.write(b"0000000000 65535 f \n")
    for offset in offsets:
        buffer.write(f"{offset:010d} 00000 n \n".encode("ascii"))
    buffer.write(f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii"))
    return buffer.getvalue()


def build_docx(paragraphs: list[tuple[bool, str]]) -> bytes:
    """A DOCX where bullets use Word's List Bullet style (no glyph in text)."""
    document = Document()
    for is_bullet, text in paragraphs:
        if is_bullet:
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def standard_pdf(sections: list[Section] | None = None) -> bytes:
    return build_pdf(pdf_lines(sections or STANDARD_SECTIONS))


def standard_docx(sections: list[Section] | None = None) -> bytes:
    return build_docx(docx_paragraphs(sections or STANDARD_SECTIONS))
