"""Integration tests for resume upload/parsing/fact-review endpoints.

Storage is redirected to a per-test `tmp_path` directory (never the real
`backend/storage/resumes/`) by monkeypatching `get_settings` as imported
into `app.api.resumes` -- this keeps the test suite from littering real
disk storage with throwaway uploads, mirroring how the rest of the suite
never makes real network calls either.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.core.config import get_settings


@pytest.fixture()
def isolated_resume_settings(tmp_path, monkeypatch):
    overridden = get_settings().model_copy(
        update={"resume_storage_dir": str(tmp_path / "resumes"), "resume_max_upload_mb": 10.0}
    )
    monkeypatch.setattr("app.api.resumes.get_settings", lambda: overridden)
    return overridden


def _build_minimal_pdf(lines: list[str]) -> bytes:
    content_ops = ["BT", "/F1 12 Tf", "72 712 Td"]
    for i, line in enumerate(lines):
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        if i > 0:
            content_ops.append("0 -14 Td")
        content_ops.append(f"({escaped}) Tj")
    content_ops.append("ET")
    stream_body = "\n".join(content_ops).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream_body)).encode("ascii") + b" >>\nstream\n" + stream_body + b"\nendstream",
    ]

    buffer = io.BytesIO()
    buffer.write(b"%PDF-1.4\n")
    offsets = [0]
    for i, body in enumerate(objects, start=1):
        offsets.append(buffer.tell())
        buffer.write(f"{i} 0 obj\n".encode("ascii"))
        buffer.write(body)
        buffer.write(b"\nendobj\n")

    xref_offset = buffer.tell()
    n = len(objects) + 1
    buffer.write(f"xref\n0 {n}\n".encode("ascii"))
    buffer.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        buffer.write(f"{off:010d} 00000 n \n".encode("ascii"))
    buffer.write(f"trailer\n<< /Size {n} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii"))
    return buffer.getvalue()


_SAMPLE_RESUME_TEXT = [
    "Jane Doe",
    "EDUCATION",
    "University of Example - B.S. Computer Science",
    "2021 - 2025",
    "SKILLS",
    "Languages: Python, TypeScript",
]


def _make_candidate(client):
    resp = client.post(
        "/api/candidate",
        json={"full_name": "Jane Doe", "email": "jane.doe@example.com"},
    )
    assert resp.status_code == 201
    return resp.json()


def test_upload_requires_existing_candidate(client, isolated_resume_settings):
    files = {"file": ("resume.pdf", _build_minimal_pdf(["Jane Doe"]), "application/pdf")}
    resp = client.post("/api/resumes", files=files)
    assert resp.status_code == 404


def test_upload_rejects_unsupported_extension(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.txt", b"plain text resume", "text/plain")}
    resp = client.post("/api/resumes", files=files)
    assert resp.status_code == 415


def test_upload_rejects_oversized_file(client, tmp_path, monkeypatch):
    _make_candidate(client)
    tiny_limit = get_settings().model_copy(
        update={"resume_storage_dir": str(tmp_path / "resumes"), "resume_max_upload_mb": 0.0001}
    )
    monkeypatch.setattr("app.api.resumes.get_settings", lambda: tiny_limit)

    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    resp = client.post("/api/resumes", files=files)
    assert resp.status_code == 413


def test_upload_pdf_success_creates_resume_and_facts(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    resp = client.post("/api/resumes", files=files)

    assert resp.status_code == 201
    body = resp.json()
    assert body["resume"]["parse_status"] == "success"
    assert body["resume"]["original_filename"] == "resume.pdf"
    assert body["facts_extracted"] > 0


def test_upload_docx_success(client, isolated_resume_settings):
    _make_candidate(client)
    document = Document()
    for line in _SAMPLE_RESUME_TEXT:
        document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)

    files = {
        "file": (
            "resume.docx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/api/resumes", files=files)
    assert resp.status_code == 201
    assert resp.json()["resume"]["parse_status"] == "success"


def test_malformed_pdf_upload_succeeds_but_parse_fails_gracefully(client, isolated_resume_settings):
    """A parse failure must never lose the original upload or crash the
    request -- the resume row is persisted with parse_status=failed."""
    _make_candidate(client)
    files = {"file": ("resume.pdf", b"%PDF-1.4\nthis is not actually a valid pdf body", "application/pdf")}
    resp = client.post("/api/resumes", files=files)

    assert resp.status_code == 201
    body = resp.json()
    assert body["resume"]["parse_status"] == "failed"
    assert body["resume"]["parse_error"]
    assert body["facts_extracted"] == 0

    # The resume row genuinely persisted, not just reflected in the response.
    get_resp = client.get(f"/api/resumes/{body['resume']['id']}")
    assert get_resp.status_code == 200
    assert get_resp.json()["parse_status"] == "failed"


def test_list_resumes_and_facts_for_resume(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    upload = client.post("/api/resumes", files=files).json()
    resume_id = upload["resume"]["id"]

    list_resp = client.get("/api/resumes")
    assert list_resp.status_code == 200
    assert any(r["id"] == resume_id for r in list_resp.json())

    facts_resp = client.get(f"/api/resumes/{resume_id}/facts")
    assert facts_resp.status_code == 200
    facts = facts_resp.json()
    assert len(facts) == upload["facts_extracted"]
    assert all(f["verification_status"] == "unverified" for f in facts)
    assert all(f["resume_id"] == resume_id for f in facts)


def test_candidate_facts_filtering_by_type(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    client.post("/api/resumes", files=files)

    resp = client.get("/api/candidate/facts", params={"fact_type": "skill"})
    assert resp.status_code == 200
    facts = resp.json()
    assert len(facts) > 0
    assert all(f["fact_type"] == "skill" for f in facts)


def test_update_fact_marks_verified(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    client.post("/api/resumes", files=files)
    fact_id = client.get("/api/candidate/facts").json()[0]["id"]

    resp = client.patch(f"/api/candidate/facts/{fact_id}", json={"verification_status": "verified"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["verification_status"] == "verified"
    assert body["verified_at"] is not None


def test_update_fact_editing_content_without_status_auto_marks_edited(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    client.post("/api/resumes", files=files)
    fact_id = client.get("/api/candidate/facts").json()[0]["id"]

    resp = client.patch(f"/api/candidate/facts/{fact_id}", json={"fact_key": "Corrected label"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["fact_key"] == "Corrected label"
    assert body["verification_status"] == "edited"


def test_update_fact_rejects_reverting_to_unverified(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    client.post("/api/resumes", files=files)
    fact_id = client.get("/api/candidate/facts").json()[0]["id"]

    resp = client.patch(f"/api/candidate/facts/{fact_id}", json={"verification_status": "unverified"})
    assert resp.status_code == 422


def test_rejected_facts_are_kept_not_deleted(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    client.post("/api/resumes", files=files)
    fact_id = client.get("/api/candidate/facts").json()[0]["id"]

    resp = client.patch(f"/api/candidate/facts/{fact_id}", json={"verification_status": "rejected"})
    assert resp.status_code == 200

    all_facts = client.get("/api/candidate/facts").json()
    assert any(f["id"] == fact_id and f["verification_status"] == "rejected" for f in all_facts)


def test_resume_has_no_generic_update_endpoint(client, isolated_resume_settings):
    """File metadata is immutable: there is no PATCH/PUT for a resume."""
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    resume_id = client.post("/api/resumes", files=files).json()["resume"]["id"]

    assert client.patch(f"/api/resumes/{resume_id}", json={"original_filename": "hacked.pdf"}).status_code == 405
    assert client.put(f"/api/resumes/{resume_id}", json={"original_filename": "hacked.pdf"}).status_code == 405


def test_upload_word_per_line_pdf_does_not_create_one_fact_per_word(client, isolated_resume_settings):
    _make_candidate(client)
    files = {
        "file": (
            "resume.pdf",
            _build_minimal_pdf(
                [
                    "Jane Doe",
                    "EDUCATION",
                    "University",
                    "of",
                    "Example",
                    "B.S.",
                    "Computer",
                    "Science",
                    "SKILLS",
                    "Languages:",
                    "Python,",
                    "TypeScript",
                ]
            ),
            "application/pdf",
        )
    }
    resp = client.post("/api/resumes", files=files)
    assert resp.status_code == 201
    facts = client.get("/api/candidate/facts").json()
    keys = {f["fact_key"] for f in facts}
    assert "of" not in keys
    assert "University" not in keys
    education = [f for f in facts if f["fact_type"] == "education"]
    assert len(education) == 1
    assert "University of Example" in education[0]["fact_key"]
    assert all(f["verification_status"] == "unverified" for f in facts)


def test_reprocess_replaces_unverified_facts_and_preserves_verified(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    upload = client.post("/api/resumes", files=files).json()
    resume_id = upload["resume"]["id"]
    original_facts = client.get(f"/api/resumes/{resume_id}/facts").json()
    assert original_facts
    keep_id = original_facts[0]["id"]
    client.patch(f"/api/candidate/facts/{keep_id}", json={"verification_status": "verified"})

    resp = client.post(f"/api/resumes/{resume_id}/reprocess")
    assert resp.status_code == 200
    body = resp.json()
    assert body["facts_preserved"] == 1
    assert body["facts_replaced"] == len(original_facts) - 1
    assert body["facts_extracted"] > 0

    after = client.get(f"/api/resumes/{resume_id}/facts").json()
    assert any(f["id"] == keep_id and f["verification_status"] == "verified" for f in after)
    # The original file is still listed -- reprocess never deletes the upload.
    listed = client.get("/api/resumes").json()
    assert any(r["id"] == resume_id for r in listed)


def test_bulk_verify_is_explicit_and_rejects_unverified_status(client, isolated_resume_settings):
    _make_candidate(client)
    files = {"file": ("resume.pdf", _build_minimal_pdf(_SAMPLE_RESUME_TEXT), "application/pdf")}
    client.post("/api/resumes", files=files)
    facts = client.get("/api/candidate/facts").json()
    ids = [f["id"] for f in facts[:2]]

    bad = client.post(
        "/api/candidate/facts/bulk",
        json={"fact_ids": ids, "verification_status": "unverified"},
    )
    assert bad.status_code == 422

    ok = client.post(
        "/api/candidate/facts/bulk",
        json={"fact_ids": ids, "verification_status": "verified"},
    )
    assert ok.status_code == 200
    assert all(f["verification_status"] == "verified" for f in ok.json())
    assert all(f["verified_at"] is not None for f in ok.json())


def test_docx_upload_with_paragraphs_and_table_is_not_word_split(client, isolated_resume_settings):
    _make_candidate(client)
    document = Document()
    document.add_paragraph("EDUCATION")
    document.add_paragraph("University of Example - B.S. Computer Science")
    document.add_paragraph("SKILLS")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "TypeScript"
    buf = io.BytesIO()
    document.save(buf)

    files = {
        "file": (
            "resume.docx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/api/resumes", files=files)
    assert resp.status_code == 201
    facts = client.get("/api/candidate/facts").json()
    keys = {f["fact_key"] for f in facts}
    assert "of" not in keys
    education = [f for f in facts if f["fact_type"] == "education"]
    assert len(education) == 1
    skills = [f for f in facts if f["fact_type"] == "skill"]
    assert skills
    items = [item for skill in skills for item in skill["fact_value"]["items"]]
    assert "Python" in items
    assert "TypeScript" in items


_SECOND_RESUME_TEXT = [
    "Jane Doe",
    "EDUCATION",
    "Example State University - B.S. Computer Science",
    "2020 - 2024",
    "SKILLS",
    "Languages: Go, Rust",
]


def _upload_resume(client, filename: str, lines: list[str]):
    files = {"file": (filename, _build_minimal_pdf(lines), "application/pdf")}
    resp = client.post("/api/resumes", files=files)
    assert resp.status_code == 201, resp.text
    return resp.json()


def test_first_upload_is_primary_and_second_is_not(client, isolated_resume_settings):
    _make_candidate(client)
    first = _upload_resume(client, "first.pdf", _SAMPLE_RESUME_TEXT)
    second = _upload_resume(client, "second.pdf", _SECOND_RESUME_TEXT)
    assert first["resume"]["is_primary"] is True
    assert second["resume"]["is_primary"] is False
    listed = {row["id"]: row for row in client.get("/api/resumes").json()}
    assert listed[first["resume"]["id"]]["is_primary"] is True
    assert listed[second["resume"]["id"]]["is_primary"] is False


def test_set_primary_resume_moves_the_flag(client, isolated_resume_settings):
    _make_candidate(client)
    first = _upload_resume(client, "first.pdf", _SAMPLE_RESUME_TEXT)["resume"]
    second = _upload_resume(client, "second.pdf", _SECOND_RESUME_TEXT)["resume"]

    resp = client.post(f"/api/resumes/{second['id']}/primary")
    assert resp.status_code == 200
    assert resp.json()["is_primary"] is True
    listed = {row["id"]: row for row in client.get("/api/resumes").json()}
    assert listed[first["id"]]["is_primary"] is False
    assert listed[second["id"]]["is_primary"] is True


def test_delete_resume_removes_only_that_files_unverified_facts(client, isolated_resume_settings):
    _make_candidate(client)
    first = _upload_resume(client, "first.pdf", _SAMPLE_RESUME_TEXT)
    second = _upload_resume(client, "second.pdf", _SECOND_RESUME_TEXT)
    first_id = first["resume"]["id"]
    second_id = second["resume"]["id"]
    first_facts = client.get(f"/api/resumes/{first_id}/facts").json()
    second_facts = client.get(f"/api/resumes/{second_id}/facts").json()
    assert first_facts and second_facts

    resp = client.request("DELETE", f"/api/resumes/{second_id}", json={})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["deleted_resume_id"] == second_id
    assert body["facts_deleted"] == len(second_facts)
    assert body["facts_preserved"] == 0
    assert body["primary_resume_id"] == first_id

    listed = client.get("/api/resumes").json()
    assert [row["id"] for row in listed] == [first_id]
    assert client.get(f"/api/resumes/{second_id}").status_code == 404
    remaining = client.get("/api/candidate/facts").json()
    remaining_ids = {fact["id"] for fact in remaining}
    assert {fact["id"] for fact in first_facts} <= remaining_ids
    assert remaining_ids.isdisjoint({fact["id"] for fact in second_facts})
    assert all(fact["resume_id"] == first_id for fact in remaining)


def test_delete_preserves_verified_facts_unless_explicitly_removed(client, isolated_resume_settings):
    _make_candidate(client)
    upload = _upload_resume(client, "only.pdf", _SAMPLE_RESUME_TEXT)
    resume_id = upload["resume"]["id"]
    facts = client.get(f"/api/resumes/{resume_id}/facts").json()
    keep = facts[0]
    client.patch(f"/api/candidate/facts/{keep['id']}", json={"verification_status": "verified"})

    resp = client.request("DELETE", f"/api/resumes/{resume_id}", json={"delete_verified_facts": False})
    assert resp.status_code == 200
    assert resp.json()["facts_preserved"] == 1
    assert resp.json()["facts_deleted"] == len(facts) - 1
    assert resp.json()["primary_resume_id"] is None

    leftover = client.get("/api/candidate/facts").json()
    assert any(fact["id"] == keep["id"] and fact["resume_id"] is None for fact in leftover)
    assert keep["id"] in {fact["id"] for fact in leftover}


def test_delete_can_remove_verified_facts_when_explicitly_chosen(client, isolated_resume_settings):
    _make_candidate(client)
    upload = _upload_resume(client, "only.pdf", _SAMPLE_RESUME_TEXT)
    resume_id = upload["resume"]["id"]
    facts = client.get(f"/api/resumes/{resume_id}/facts").json()
    keep = facts[0]
    client.patch(f"/api/candidate/facts/{keep['id']}", json={"verification_status": "verified"})

    resp = client.request("DELETE", f"/api/resumes/{resume_id}", json={"delete_verified_facts": True})
    assert resp.status_code == 200
    assert resp.json()["facts_preserved"] == 0
    leftover = client.get("/api/candidate/facts").json()
    assert keep["id"] not in {fact["id"] for fact in leftover}


def test_delete_primary_requires_replacement_when_others_exist(client, isolated_resume_settings):
    _make_candidate(client)
    first = _upload_resume(client, "first.pdf", _SAMPLE_RESUME_TEXT)["resume"]
    second = _upload_resume(client, "second.pdf", _SECOND_RESUME_TEXT)["resume"]

    missing = client.request("DELETE", f"/api/resumes/{first['id']}", json={})
    assert missing.status_code == 409
    assert client.get(f"/api/resumes/{first['id']}").status_code == 200

    ok = client.request(
        "DELETE",
        f"/api/resumes/{first['id']}",
        json={"replacement_resume_id": second["id"]},
    )
    assert ok.status_code == 200, ok.text
    assert ok.json()["primary_resume_id"] == second["id"]
    listed = client.get("/api/resumes").json()
    assert len(listed) == 1
    assert listed[0]["id"] == second["id"]
    assert listed[0]["is_primary"] is True


def test_delete_unknown_resume_returns_404(client, isolated_resume_settings):
    _make_candidate(client)
    resp = client.request("DELETE", "/api/resumes/999999", json={})
    assert resp.status_code == 404


def test_delete_resume_does_not_affect_jobs_preferences_or_saved_jobs(client, isolated_resume_settings, monkeypatch):
    _make_candidate(client)
    prefs_payload = {
        "remote_ok": True,
        "target_employment_types": ["internship"],
        "target_start_window_start": None,
        "target_start_window_end": None,
        "min_match_score_to_surface": 10,
        "locations": [{"label": "Remote", "location_type": "remote", "priority": 0}],
        "roles": [{"role_category": "swe_backend", "priority": 0}],
    }
    prefs = client.put("/api/candidate/preferences", json=prefs_payload)
    assert prefs.status_code == 200, prefs.text

    import json
    from pathlib import Path

    fixtures = Path(__file__).resolve().parents[1] / "fixtures"

    class _FakeResponse:
        def __init__(self, status_code: int, json_data):
            self.status_code = status_code
            self._json_data = json_data
            self.request = None

        def json(self):
            return self._json_data

    monkeypatch.setattr(
        "app.sources.http.httpx.get",
        lambda url, params=None, timeout=None: _FakeResponse(
            200, json.loads((fixtures / "greenhouse_jobs.json").read_text())
        ),
    )
    client.post(
        "/api/job-sources",
        json={"employer_name": "Acme Corp", "provider": "greenhouse", "board_token": "acme"},
    )
    run = client.post("/api/discovery/runs")
    assert run.status_code == 201, run.text
    jobs = client.get("/api/jobs").json()
    assert jobs
    job_id = jobs[0]["id"]
    saved = client.patch(f"/api/jobs/{job_id}/state", json={"state": "saved"})
    assert saved.status_code == 200

    resume = _upload_resume(client, "only.pdf", _SAMPLE_RESUME_TEXT)["resume"]
    deleted = client.request("DELETE", f"/api/resumes/{resume['id']}", json={})
    assert deleted.status_code == 200

    after_prefs = client.get("/api/candidate/preferences").json()
    assert after_prefs["min_match_score_to_surface"] == 10
    assert [loc["label"] for loc in after_prefs["locations"]] == ["Remote"]
    after_jobs = client.get("/api/jobs").json()
    assert len(after_jobs) == len(jobs)
    saved_list = client.get("/api/jobs", params={"state": "saved"}).json()
    assert len(saved_list) == 1
    assert saved_list[0]["id"] == job_id
