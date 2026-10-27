"""Format-specific text extraction. The only job of this module is "get the
plain text out of the file bytes" -- no section/fact logic lives here, so
it stays swappable if a format gains a better extraction library later.
"""

from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.exceptions import ResumeExtractionError
from app.models.enums import ResumeFileFormat
from app.resume_parsing.text_normalize import fragmentation_score, normalize_extracted_text


def extract_text(file_bytes: bytes, file_format: ResumeFileFormat) -> str:
    """Return the plain text content of a resume file.

    Raises `ResumeExtractionError` if the bytes can't be read as the given
    format (corrupt file, encrypted PDF, etc.) or if extraction succeeds
    but yields no usable text at all. Never returns fabricated content.
    """
    if file_format == ResumeFileFormat.PDF:
        return _extract_pdf_text(file_bytes)
    if file_format == ResumeFileFormat.DOCX:
        return _extract_docx_text(file_bytes)
    # Unreachable given the ResumeFileFormat enum, but fail loudly rather
    # than silently returning empty text if a new format is ever added
    # here without updating this function.
    raise ResumeExtractionError(f"No extractor implemented for format: {file_format}")


def _choose_pdf_page_text(page) -> str:
    """Prefer pypdf layout mode: Google Docs / Word PDFs store each word
    as its own positioned run, and plain mode emits one token per line.
    Fall back to whichever candidate is less fragmented."""
    try:
        layout = page.extract_text(extraction_mode="layout") or ""
    except Exception:
        layout = ""
    try:
        plain = page.extract_text() or ""
    except Exception:
        plain = ""

    layout = layout.strip()
    plain = plain.strip()
    if layout and plain:
        return layout if fragmentation_score(layout) <= fragmentation_score(plain) else plain
    return layout or plain


def _extract_pdf_text(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted:
            raise ResumeExtractionError("PDF is password-protected; cannot extract text.")
        pages_text = [_choose_pdf_page_text(page) for page in reader.pages]
    except PdfReadError as exc:
        raise ResumeExtractionError(f"Could not read PDF: {exc}") from exc
    except ResumeExtractionError:
        raise
    except Exception as exc:  # pypdf can raise a variety of error types on malformed input
        raise ResumeExtractionError(f"Could not read PDF: {exc}") from exc

    text = normalize_extracted_text("\n".join(pages_text))
    if not text:
        raise ResumeExtractionError("PDF contained no extractable text (scanned image? empty file?).")
    return text


def _is_docx_list_item(paragraph) -> bool:
    """True for real Word list items (bullets or numbering)."""
    try:
        if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
            return True
    except AttributeError:
        pass
    style = getattr(getattr(paragraph, "style", None), "name", "") or ""
    return style.lower().startswith("list ") and style.lower() != "list paragraph"


def _docx_paragraph_text(paragraph) -> str:
    """Paragraph text, with a bullet glyph for Word list items.

    Word stores the bullet as list formatting, not as text, so without this a
    DOCX bullet is indistinguishable from a heading or a company line.
    """
    text = paragraph.text or ""
    if text.strip() and _is_docx_list_item(paragraph):
        return f"• {text.strip()}"
    return text


def _extract_docx_text(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))
        paragraphs = [_docx_paragraph_text(p) for p in document.paragraphs]
        # Resumes sometimes lay out skills/experience in tables rather than
        # paragraphs. Join a row's cells onto one line so each cell does
        # not become its own "fact" later. Content is never dropped.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
    except Exception as exc:  # python-docx raises plain Exception/KeyError/zipfile errors on bad input
        raise ResumeExtractionError(f"Could not read DOCX: {exc}") from exc

    text = normalize_extracted_text("\n".join(p for p in paragraphs if p is not None))
    if not text:
        raise ResumeExtractionError("DOCX contained no extractable text.")
    return text
