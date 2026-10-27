"""ATS-friendly DOCX/PDF export of a tailored resume document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from fpdf import FPDF

from app.services.resume_tailor import flatten_tailored_lines, structured_entry_lines


def export_docx(document: dict, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    header = document.get("header") or {}
    name = header.get("full_name") or "Candidate"
    doc.add_heading(name, level=0)
    contact = " · ".join(part for part in (header.get("email"), header.get("phone"), header.get("location")) if part)
    if contact:
        doc.add_paragraph(contact)
    for section in ("education", "work_experience", "projects", "skills", "achievements"):
        rows = document.get(section) or []
        if not rows:
            continue
        doc.add_heading(section.replace("_", " ").title(), level=1)
        for row in rows:
            lines = structured_entry_lines(row, section)
            if not lines:
                continue
            heading = doc.add_paragraph()
            run = heading.add_run(lines[0])
            run.bold = True
            for line in lines[1:]:
                if line.startswith("• "):
                    doc.add_paragraph(line[2:], style="List Bullet")
                else:
                    doc.add_paragraph(line)
    doc.save(destination)
    return destination


def export_pdf(document: dict, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in flatten_tailored_lines(document):
        pdf.multi_cell(pdf.epw, 8, text=line.encode("latin-1", "replace").decode("latin-1"))
    pdf.output(destination)
    return destination
